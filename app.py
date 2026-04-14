from flask import Flask, render_template, request, jsonify, send_file
import pdfplumber
import re
import io
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

app = Flask(__name__)
app.config['MAX_CONTENT_LENGTH'] = 100 * 1024 * 1024


def extract_data_from_pdf(pdf_bytes, filename=''):
    results = []
    with pdfplumber.open(io.BytesIO(pdf_bytes)) as pdf:
        for page in pdf.pages:
            text = page.extract_text(x_tolerance=3, y_tolerance=3) or ""
            if not text:
                continue
            lines = [l.strip() for l in text.split('\n')]

            # Find tracking number on this page
            tracking_no = None
            for line in lines:
                m = re.search(r'\b(1Z[A-Z0-9]{16})\b', line, re.IGNORECASE)
                if m:
                    tracking_no = m.group(1).upper()
                    break
            if not tracking_no:
                continue

            # Parse net charges section
            in_charges = False
            for line in lines:
                if not line:
                    continue
                if re.search(r'\bNet\s*Charges\b', line, re.IGNORECASE):
                    in_charges = True
                    continue
                if in_charges and re.match(r'^(Total Charges|Date\s|Page \d|Tracking)', line, re.IGNORECASE):
                    in_charges = False
                if not in_charges:
                    continue
                if line.lower() in ['description', 'net charges']:
                    continue

                # Match "Some Description   123.45"
                m = re.match(r'^(.+?)\s+([\d,]+\.\d{2})\s*$', line)
                if m:
                    desc = m.group(1).strip()
                    skip = ['description', 'shipper', 'consignee', 'payor', 'comments', 'reference']
                    if any(s in desc.lower() for s in skip):
                        continue
                    try:
                        amount = float(m.group(2).replace(',', ''))
                        if amount > 0:
                            results.append({
                                'tracking_no': tracking_no,
                                'description': desc,
                                'net_charges': round(amount, 2)
                            })
                    except ValueError:
                        pass

            # Fallback: known charge keywords
            if not any(r['tracking_no'] == tracking_no for r in results):
                known = ['Duty Amount', 'Value Added Tax', 'Import Tax', 'Customs Duty',
                         'Fuel Surcharge', 'Peak Surcharge', 'Residential Surcharge',
                         'Disbursement Fee', 'Handling Fee', 'Transportation Charge']
                for charge in known:
                    m = re.search(rf'{re.escape(charge)}\s+([\d,]+\.\d{{2}})', text, re.IGNORECASE)
                    if m:
                        try:
                            amount = float(m.group(1).replace(',', ''))
                            if amount > 0:
                                results.append({
                                    'tracking_no': tracking_no,
                                    'description': charge,
                                    'net_charges': round(amount, 2)
                                })
                        except ValueError:
                            pass

    # Deduplicate
    seen = set()
    unique = []
    for r in results:
        key = (r['tracking_no'], r['description'].lower(), r['net_charges'])
        if key not in seen:
            seen.add(key)
            unique.append(r)
    return unique


@app.route('/')
def index():
    return render_template('index.html')


@app.route('/upload', methods=['POST'])
def upload():
    files = request.files.getlist('files')
    if not files:
        return jsonify({'error': 'No files', 'data': [], 'total': 0}), 400

    all_data, errors = [], []
    for f in files:
        if not f.filename.lower().endswith('.pdf'):
            errors.append(f'{f.filename}: Not a PDF')
            continue
        try:
            data = extract_data_from_pdf(f.read(), f.filename)
            all_data.extend(data)
        except Exception as e:
            errors.append(f'{f.filename}: {str(e)}')

    return jsonify({'data': all_data, 'total': len(all_data), 'errors': errors})


@app.route('/download', methods=['POST'])
def download():
    data = request.json.get('data', [])
    if not data:
        return jsonify({'error': 'No data'}), 400

    wb = Workbook()
    ws = wb.active
    ws.title = "UPS Invoice Data"

    hfill = PatternFill(start_color="351C75", end_color="351C75", fill_type="solid")
    hfont = Font(bold=True, color="FFFFFF", size=11)
    afill = PatternFill(start_color="F3F0FF", end_color="F3F0FF", fill_type="solid")
    border = Border(*[Side(style='thin', color='CCCCCC')] * 0,
                    left=Side(style='thin'), right=Side(style='thin'),
                    top=Side(style='thin'), bottom=Side(style='thin'))

    headers = ['Tracking No.', 'Description', 'Net Charges']
    ws.row_dimensions[1].height = 32
    for col, h in enumerate(headers, 1):
        c = ws.cell(row=1, column=col, value=h)
        c.font = hfont
        c.fill = hfill
        c.alignment = Alignment(horizontal='center', vertical='center')
        c.border = border

    for i, item in enumerate(data, 2):
        ws.cell(row=i, column=1, value=item['tracking_no']).border = border
        ws.cell(row=i, column=2, value=item['description']).border = border
        nc = ws.cell(row=i, column=3, value=item['net_charges'])
        nc.number_format = '#,##0.00'
        nc.alignment = Alignment(horizontal='right')
        nc.border = border
        if i % 2 == 0:
            for col in range(1, 4):
                ws.cell(row=i, column=col).fill = afill

    ws.column_dimensions['A'].width = 26
    ws.column_dimensions['B'].width = 38
    ws.column_dimensions['C'].width = 15
    ws.freeze_panes = 'A2'

    out = io.BytesIO()
    wb.save(out)
    out.seek(0)
    return send_file(out,
                     mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
                     as_attachment=True, download_name='ups_invoice_data.xlsx')


if __name__ == '__main__':
    app.run(host='0.0.0.0', port=5002, debug=False)
